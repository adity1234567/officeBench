"""Collect every gemma4:31b failure (plus the one false pass) into analysis/failures.json."""
import json, os, io, contextlib, glob
import openpyxl, icalendar
from utils import evaluate as ev

CASES = [("1-2/0","first10"),("1-2/1","first10"),("1-2/2","first10"),("1-2/4","first10"),
         ("1-9/1","probe"),("3-82/0","probe"),
         ("1-3/4","next20"),("1-4/2","next20"),("1-4/3","next20"),("1-5/1","next20"),("1-5/2","next20"),
         ("1-5/3","next20"),("1-6/0","next20"),("1-6/2","next20"),("1-6/3","next20"),("1-6/4","next20")]

def dump_state(testbed):
    out = {}
    for p in sorted(glob.glob(f'{testbed}/**/*', recursive=True)):
        rel = os.path.relpath(p, testbed).replace(os.sep, '/')
        if os.path.isdir(p): continue
        if p.endswith('.xlsx'):
            ws = openpyxl.load_workbook(p).active
            out[rel] = [[c.value for c in r] for r in ws.iter_rows()]
        elif p.endswith('.ics'):
            cal = icalendar.Calendar.from_ical(open(p, 'rb').read())
            out[rel] = [f"{e.get('summary')} {e.get('dtstart').dt}–{e.get('dtend').dt}" for e in cal.walk('VEVENT')]
        elif p.endswith('.txt'):
            out[rel] = open(p, encoding='utf-8', errors='replace').read()[:400]
        elif p.endswith('.docx'):
            import docx
            out[rel] = '\n'.join(par.text for par in docx.Document(p).paragraphs)[:400]
        else:
            out[rel] = '(file present)'
    return out

cases = []
for tid, tag in CASES:
    t, s = tid.split('/')
    cfg = json.load(open(f'tasks/{t}/subtasks/{s}.json', encoding='utf-8'))
    d = f'tasks/{t}/outputs/{s}/gemma4_31b_{tag}'
    hist = json.load(open(f'{d}/env_history.json', encoding='utf-8'))
    checks = []
    for e in cfg['evaluation']:
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            try: ok = getattr(ev, e['function'])(f'{d}/testbed', e['args'])
            except Exception as ex: ok = f'ERROR: {ex}'
        checks.append({'function': e['function'], 'args': e['args'], 'result': ok})
    cases.append({
        'id': tid, 'run_tag': tag, 'task': cfg['task'],
        'scored': 'pass' if all(c['result'] is True for c in checks) else 'fail',
        'checks': checks,
        'steps': [{'action': str(a)[:400], 'observation': str(o)[:400]} for a, o in hist],
        'initial_state': dump_state(f'tasks/{t}/testbed'),
        'final_state': dump_state(f'{d}/testbed'),
    })
json.dump(cases, open('analysis/failures.json', 'w', encoding='utf-8'), indent=1, ensure_ascii=False, default=str)
print(len(cases), 'cases;', [(c['id'], c['scored']) for c in cases])
